import csv
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from fastapi import HTTPException
from fastapi.responses import FileResponse, RedirectResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .filing_availability import check_company_filing
from .index_registry import DASHBOARD_DATA_DIR, ROOT, companies_for_index


EXPORT_DIR = ROOT / "exports"
JOB_DIR = EXPORT_DIR / "jobs"


def _load_company(symbol: str) -> Dict[str, Any]:
    path = DASHBOARD_DATA_DIR / f"{symbol.upper()}.json"
    if not path.exists():
        raise HTTPException(status_code=404, detail=f"Parsed financial data not found for {symbol.upper()}")
    return json.loads(path.read_text(encoding="utf-8"))


def _safe_sheet_name(name: str) -> str:
    return "".join(ch for ch in name if ch not in "[]:*?/\\")[:31] or "Sheet"


def _sorted_periods(data: Dict[str, Any]) -> List[str]:
    def key(period: str) -> tuple:
        text = str(period)
        if text.startswith("FY") and text[2:].isdigit():
            return (int(text[2:]), 12)
        try:
            dt = datetime.strptime(text, "%b %Y")
            return (dt.year, dt.month)
        except ValueError:
            return (0, 0)

    return sorted((data or {}).keys(), key=key, reverse=True)


def _append_mapping_sheet(wb: Workbook, title: str, rows: Dict[str, Any]) -> None:
    ws = wb.create_sheet(_safe_sheet_name(title))
    ws.append(["Key", "Value"])
    for key, value in rows.items():
        ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    _style_sheet(ws)


def _append_statement_sheet(wb: Workbook, title: str, statement: Dict[str, Dict[str, Any]]) -> None:
    ws = wb.create_sheet(_safe_sheet_name(title))
    periods = _sorted_periods(statement)
    fields = sorted({field for period in periods for field in (statement.get(period) or {}).keys()})
    ws.append(["Reported Item", *periods])
    for field in fields:
        ws.append([field, *[(statement.get(period) or {}).get(field) for period in periods]])
    _style_sheet(ws)


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 42)


def build_company_workbook(symbol: str, period_end: str = "2026-03-31", result_type: str = "all") -> bytes:
    data = _load_company(symbol)
    company = data.get("company", {})
    requested_type = result_type if result_type in ("consolidated", "standalone") else "all"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    availability = check_company_filing(symbol, period_end)
    summary = {
        "Company": company.get("name"),
        "Symbol": company.get("symbol") or symbol.upper(),
        "Sector": company.get("sector"),
        "Industry": company.get("industry"),
        "Period End": period_end,
        "Result Type": "All available consolidated and standalone filings" if requested_type == "all" else requested_type,
        "Filing Status": availability.get("status"),
        "Generated At": datetime.now(timezone.utc).isoformat(),
        "Unit": data.get("metadata", {}).get("unit", "₹ Crores"),
    }
    ws.append(["Key", "Value"])
    for key, value in summary.items():
        ws.append([key, value])
    _style_sheet(ws)

    financials = data.get("financials") or {}
    result_types = ["consolidated", "standalone"] if requested_type == "all" else [requested_type]
    for current_type in result_types:
        financial_layer = financials.get(current_type) or {}
        for period_name, statements in financial_layer.items():
            for statement_name, statement in (statements or {}).items():
                if statement:
                    sheet = f"{current_type[:4]}_{period_name[:3]}_{statement_name}"
                    _append_statement_sheet(wb, sheet, statement)

    if data.get("ratios"):
        _append_statement_sheet(wb, "Ratios", data["ratios"])
    elif data.get("derived_metrics"):
        _append_statement_sheet(wb, "Derived Metrics", data["derived_metrics"])
    if data.get("graph_data"):
        _append_mapping_sheet(wb, "Graph Data", data["graph_data"])
    _append_mapping_sheet(wb, "Metadata", data.get("metadata", {}))

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def company_export_response(symbol: str, fmt: str, period_end: str = "2026-03-31", result_type: str = "consolidated") -> Response:
    fmt = fmt.lower()
    symbol = symbol.upper()
    if fmt == "xlsx":
        payload = build_company_workbook(symbol, period_end, result_type)
        return Response(
            payload,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{symbol}_Financials_{period_end}.xlsx"'},
        )
    if fmt == "json":
        data = _load_company(symbol)
        return Response(
            json.dumps(data, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{symbol}_Financials_{period_end}.json"'},
        )
    if fmt == "csv":
        data = _load_company(symbol)
        requested_type = result_type if result_type in ("consolidated", "standalone") else "all"
        rows = []
        financials = data.get("financials") or {}
        result_types = ["consolidated", "standalone"] if requested_type == "all" else [requested_type]
        for current_type in result_types:
            financial_layer = financials.get(current_type) or {}
            for period_name, statements in financial_layer.items():
                for statement_name, statement in (statements or {}).items():
                    for period, values in (statement or {}).items():
                        for field, value in (values or {}).items():
                            rows.append(
                                {
                                    "symbol": symbol,
                                    "result_type": current_type,
                                    "period_type": period_name,
                                    "statement": statement_name,
                                    "period": period,
                                    "field": field,
                                    "value": value,
                                }
                            )
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=["symbol", "result_type", "period_type", "statement", "period", "field", "value"])
        writer.writeheader()
        writer.writerows(rows)
        return Response(
            stream.getvalue(),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{symbol}_Financials_{period_end}.csv"'},
        )

    availability = check_company_filing(symbol, period_end)
    if fmt == "xbrl" and availability.get("raw_xbrl"):
        raw_xbrl = availability["raw_xbrl"]
        if str(raw_xbrl).startswith(("http://", "https://")):
            return RedirectResponse(raw_xbrl)
        return FileResponse(raw_xbrl, filename=f"{symbol}_{period_end}.xml")
    if fmt == "pdf" and availability.get("pdf"):
        pdf = availability["pdf"]
        if str(pdf).startswith(("http://", "https://")):
            return RedirectResponse(pdf)
        return FileResponse(pdf, filename=f"{symbol}_{period_end}.pdf")
    raise HTTPException(status_code=404, detail=f"{fmt.upper()} filing asset is not available for {symbol}")


def create_bulk_excel_export(index_code: str, period_end: str = "2026-03-31") -> Dict[str, Any]:
    index_code = index_code.upper()
    JOB_DIR.mkdir(parents=True, exist_ok=True)
    job_id = f"{index_code}_{period_end}_{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}"
    job_path = JOB_DIR / job_id
    job_path.mkdir(parents=True, exist_ok=True)
    zip_path = job_path / f"{index_code}_Financials_{period_end}.zip"

    companies = companies_for_index(index_code)
    manifest = {
        "job_id": job_id,
        "index_code": index_code,
        "period_end": period_end,
        "format": "xlsx",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "companies": [],
    }

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for company in companies:
            symbol = company["symbol"]
            availability = check_company_filing(symbol, period_end)
            item = {"symbol": symbol, "name": company.get("name"), "status": availability["status"], "file": None}
            if availability["status"] == "PARSED":
                filename = f"{symbol}.xlsx"
                zf.writestr(filename, build_company_workbook(symbol, period_end, "all"))
                item["file"] = filename
            manifest["companies"].append(item)
        zf.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        zf.writestr("export_summary.csv", _summary_csv(manifest["companies"]))

    status_path = job_path / "status.json"
    status_path.write_text(
        json.dumps(
            {
                "job_id": job_id,
                "status": "completed",
                "zip_path": str(zip_path),
                "download_url": f"/api/admin/exports/{job_id}/download",
                "company_count": len(companies),
                "exported_count": sum(1 for c in manifest["companies"] if c["file"]),
                "missing_count": sum(1 for c in manifest["companies"] if not c["file"]),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return json.loads(status_path.read_text(encoding="utf-8"))


def _summary_csv(rows: Iterable[Dict[str, Any]]) -> str:
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=["symbol", "name", "status", "file"])
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def get_export_job(job_id: str) -> Dict[str, Any]:
    status_path = JOB_DIR / job_id / "status.json"
    if not status_path.exists():
        raise HTTPException(status_code=404, detail=f"Export job {job_id} not found")
    return json.loads(status_path.read_text(encoding="utf-8"))
